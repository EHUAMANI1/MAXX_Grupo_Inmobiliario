from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


# =========================
# RUTAS DE ENTRADA Y SALIDA
# =========================

BASE_DIR = Path(__file__).resolve().parent.parent

ENTRADA_DATOS = BASE_DIR / "Flujo" / "Input" / "Ventas (3) (1).xlsx"
SALIDA_DATOS = BASE_DIR / "Flujo" / "Output" / "Competencia_ventas.csv"


# =========================
# FUNCIÓN PARA LIMPIAR VALORES
# =========================

def limpiar_valor(valor):
    """
    Convierte valores vacíos en texto vacío.
    """
    if valor is None:
        return ""

    return str(valor).strip()


def corregir_coordenada(valor):
    """
    Corrige latitud y longitud:
    - cambia coma decimal por punto decimal
    - evita que Excel/Python lo interprete como miles o millones
    """
    if valor is None:
        return ""

    valor = str(valor).strip()

    if valor == "":
        return ""

    # Cambiar coma decimal por punto decimal
    valor = valor.replace(",", ".")

    return valor


# =========================
# FUNCIÓN PRINCIPAL
# =========================

def arreglar_celdas_unidas_excel(ruta_entrada, ruta_salida, nombre_hoja=None):
    """
    Lee un Excel con celdas unidas, quita las uniones,
    rellena los espacios vacíos con el valor original
    y exporta la base limpia a CSV.
    """

    if not ruta_entrada.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {ruta_entrada}")

    ruta_salida.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(ruta_entrada, data_only=True)

    if nombre_hoja is None:
        ws = wb.active
    else:
        ws = wb[nombre_hoja]

    # =========================
    # 1. QUITAR CELDAS UNIDAS
    # =========================

    rangos_combinados = list(ws.merged_cells.ranges)

    for rango in rangos_combinados:
        min_col = rango.min_col
        max_col = rango.max_col
        min_row = rango.min_row
        max_row = rango.max_row

        valor = ws.cell(row=min_row, column=min_col).value

        ws.unmerge_cells(str(rango))

        for fila in range(min_row, max_row + 1):
            for columna in range(min_col, max_col + 1):
                ws.cell(row=fila, column=columna).value = valor

    # =========================
    # 2. PASAR EXCEL A DATAFRAME
    # =========================

    datos = []

    for fila in ws.iter_rows(values_only=True):
        datos.append([limpiar_valor(valor) for valor in fila])

    df = pd.DataFrame(datos)

    # Primera fila como encabezado
    df.columns = df.iloc[0]
    df = df.iloc[1:].reset_index(drop=True)

    # Eliminar filas y columnas completamente vacías
    df = df.replace("", pd.NA)
    df = df.dropna(how="all")
    df = df.dropna(axis=1, how="all")
    df = df.fillna("")

    # =========================
    # 3. CORREGIR LATITUD Y LONGITUD
    # =========================

    columnas_coordenadas = ["Latitud", "Longitud"]

    for columna in columnas_coordenadas:
        if columna in df.columns:
            df[columna] = df[columna].apply(corregir_coordenada)

    # =========================
    # 4. EXPORTAR A CSV
    # =========================

    df.to_csv(
        ruta_salida,
        index=False,
        encoding="utf-8-sig",
        sep=";"
    )

    print("Proceso terminado correctamente.")
    print(f"Archivo exportado en: {ruta_salida}")


# =========================
# EJECUCIÓN
# =========================

if __name__ == "__main__":
    arreglar_celdas_unidas_excel(
        ruta_entrada=ENTRADA_DATOS,
        ruta_salida=SALIDA_DATOS
    )